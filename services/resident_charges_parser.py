"""
Parser for the Entrata Resident Charges multi-sheet Excel report.

Each property sheet (except "Report Parameters") has the layout:

  Row 1-4:  Title block  (blank, "Resident Charges", property name, period)
  Row 5:    Blank
  Row 6:    Column headers: Bldg-Unit | Resident | Lease Status | Amount
  Row 7+:   Data rows
  Last row: Total row  (column C = "Total:", column D = SUM formula)

Usage::

    from services.resident_charges_parser import parse_resident_charges

    records, period = parse_resident_charges("August Resident Charges.xlsx")
"""
from __future__ import annotations

import logging
import re
from decimal import Decimal
from pathlib import Path
from typing import List, Optional, Tuple

from models.resident_charges_models import ResidentChargeRecord
from services.utils import normalize_text

logger = logging.getLogger(__name__)

_SKIP_SHEETS = {"report parameters"}
_HEADER_ROW_MARKER = "bldg-unit"   # normalized value of the first header cell


def _normalize_resident_name(raw_name: str) -> str:
    """Convert a "Last, First" Entrata name to a normalized "first last" form.

    Also handles parenthetical suffixes such as "(Lewis Walker)" by
    stripping them before conversion.
    """
    if not raw_name:
        return ""
    # Remove anything inside parentheses
    name = re.sub(r"\([^)]*\)", "", raw_name).strip()
    if "," in name:
        parts = name.split(",", 1)
        last = parts[0].strip()
        first = parts[1].strip()
        name = f"{first} {last}"
    return normalize_text(name)


def parse_resident_charges(
    file_path: str,
) -> Tuple[List[ResidentChargeRecord], str]:
    """Parse every property sheet in the Resident Charges workbook.

    Args:
        file_path: Absolute or relative path to the XLSX file.

    Returns:
        Tuple of:
          - List of ``ResidentChargeRecord`` (one per resident row).
          - Reporting period string extracted from the first property sheet
            (e.g. ``"Aug 2026"``), or empty string if not found.
    """
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError(
            "openpyxl is required to parse the Resident Charges report.  "
            "Run: pip install openpyxl"
        ) from exc

    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Resident Charges file not found: {file_path}")

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    records: List[ResidentChargeRecord] = []
    reporting_period = ""

    for sheet_name in wb.sheetnames:
        if normalize_text(sheet_name) in _SKIP_SHEETS:
            logger.debug("Skipping sheet: %s", sheet_name)
            continue

        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))

        if len(rows) < 7:
            logger.warning("Sheet %r has fewer than 7 rows – skipped.", sheet_name)
            continue

        # Row 4 (index 3) contains the reporting period, e.g. "Aug 2026"
        period_cell = rows[3][0] if rows[3] else None
        if period_cell and isinstance(period_cell, str) and not reporting_period:
            reporting_period = period_cell.strip()

        # Find the header row by looking for "Bldg-Unit" in column A
        header_row_idx = None
        for i, row in enumerate(rows):
            if row and row[0] is not None:
                if normalize_text(str(row[0])) == _HEADER_ROW_MARKER:
                    header_row_idx = i
                    break

        if header_row_idx is None:
            logger.warning("Header row not found in sheet %r – skipped.", sheet_name)
            continue

        period_str = reporting_period or ""

        for row in rows[header_row_idx + 1 :]:
            if not row or row[0] is None:
                continue

            unit_val = row[0]
            resident_val = row[1]
            lease_status_val = row[2]
            amount_val = row[3]

            # Skip the total row
            if (
                lease_status_val is not None
                and normalize_text(str(lease_status_val)) == "total:"
            ):
                continue

            # Skip rows without a resident name
            if resident_val is None:
                continue

            unit_str = str(unit_val).strip() if unit_val is not None else ""
            resident_str = str(resident_val).strip()
            lease_str = str(lease_status_val).strip() if lease_status_val is not None else ""

            try:
                amount = Decimal(str(amount_val)) if amount_val is not None else Decimal("0")
            except Exception:
                amount = Decimal("0")

            records.append(
                ResidentChargeRecord(
                    property_name=sheet_name,
                    unit=unit_str,
                    resident_name=resident_str,
                    normalized_name=_normalize_resident_name(resident_str),
                    lease_status=lease_str,
                    amount=amount,
                    reporting_period=period_str,
                )
            )
            logger.debug(
                "Parsed resident: %s / %s / %s / %s",
                sheet_name,
                unit_str,
                resident_str,
                amount,
            )

    wb.close()
    logger.info(
        "Parsed %d resident charge records across %d sheets.",
        len(records),
        len(set(r.property_name for r in records)),
    )
    return records, reporting_period
