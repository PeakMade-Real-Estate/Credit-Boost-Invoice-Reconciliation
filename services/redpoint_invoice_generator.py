"""Generate Redpoint invoices from Boom statement exports."""
from __future__ import annotations

import logging
import os
import re
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from typing import Optional, Tuple

import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from services.utils import parse_decimal

logger = logging.getLogger(__name__)

_PROJECT_ROOT = Path(__file__).parent.parent
_REDPOINT_LOGO_PATH = _PROJECT_ROOT / "Redpoint_logo.png"
_CREDIT_BOOST_LOGO_PATH = _PROJECT_ROOT / "Credit Boost Logo transparent.png"
_VISIBLE_SHEET = "Redpoint Invoice"
_SUMMARY_SHEET = "Property Summary"
_RECONCILIATION_SHEET = "Reconciliation Data"

_AMOUNT_COLUMN_CANDIDATES = {"amount", "transactionamount", "normalizedamount"}
_DROP_COLUMN_KEYS = {
    "id",
    "item",
    "category",
    "description",
    "templatename",
    "template",
    "propertyid",
    "transaction",
    "transactiontype",
    "subjecttype",
    "bankaccountname",
    "bankaccountlast4",
    "statementid",
}


def generate_redpoint_invoice(
    boom_statement_path: str,
    output_folder: str,
    run_id: str,
    *,
    base_price: Optional[Decimal] = None,
) -> str:
    """Create a branded Redpoint invoice workbook from a Boom statement export."""
    xlsx_path, _pdf_path = _generate_redpoint_invoice_artifacts(
        boom_statement_path,
        output_folder,
        run_id,
        base_price=base_price,
        include_pdf=False,
    )
    return xlsx_path


def generate_redpoint_invoice_package(
    boom_statement_path: str,
    output_folder: str,
    run_id: str,
    *,
    base_price: Optional[Decimal] = None,
) -> Tuple[str, str]:
    """Create the branded XLSX invoice and matching PDF invoice."""
    xlsx_path, pdf_path = _generate_redpoint_invoice_artifacts(
        boom_statement_path,
        output_folder,
        run_id,
        base_price=base_price,
        include_pdf=True,
    )
    if not pdf_path:
        raise RuntimeError("PDF invoice generation did not return a path.")
    return xlsx_path, pdf_path


def _generate_redpoint_invoice_artifacts(
    boom_statement_path: str,
    output_folder: str,
    run_id: str,
    *,
    base_price: Optional[Decimal],
    include_pdf: bool,
) -> Tuple[str, Optional[str]]:
    source_path = Path(boom_statement_path)
    if not source_path.exists():
        raise FileNotFoundError(f"Boom statement file not found: {boom_statement_path}")

    ext = source_path.suffix.lower()
    if ext == ".csv":
        df = pd.read_csv(source_path, dtype=str, keep_default_na=False)
    elif ext == ".xlsx":
        df = pd.read_excel(source_path, dtype=str, keep_default_na=False)
    else:
        raise ValueError(f"Unsupported Boom statement file type: {ext!r}")

    price = _resolve_base_price(base_price)
    amount_col = _find_amount_column(df)
    df[amount_col] = f"{price:.2f}"

    drop_columns = [c for c in df.columns if _column_key(c) in _DROP_COLUMN_KEYS]
    invoice_df = df.drop(columns=drop_columns)
    summary_df = _build_property_summary(invoice_df)

    os.makedirs(output_folder, exist_ok=True)
    timestamp = datetime.now(UTC).strftime("%Y%m%d_%H%M%S")
    output_path = Path(output_folder) / f"{run_id}_redpoint_invoice_{timestamp}.xlsx"
    _write_invoice_workbook(output_path, invoice_df, summary_df, df, run_id)

    pdf_path = None
    if include_pdf:
        pdf_path = Path(output_folder) / f"{run_id}_redpoint_invoice_{timestamp}.pdf"
        _write_invoice_pdf(pdf_path, summary_df, run_id)

    logger.info(
        "Redpoint invoice generated from %s: %s  rows=%d  base_price=%s",
        source_path.name,
        output_path,
        len(invoice_df),
        price,
    )
    return str(output_path), str(pdf_path) if pdf_path else None


def _write_invoice_workbook(
    output_path: Path,
    invoice_df: pd.DataFrame,
    summary_df: pd.DataFrame,
    reconciliation_df: pd.DataFrame,
    run_id: str,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = _VISIBLE_SHEET

    _add_logo(ws, _REDPOINT_LOGO_PATH, "A1", width=260)
    _add_logo(ws, _CREDIT_BOOST_LOGO_PATH, "D1", width=210)

    ws["A5"] = "Invoice"
    ws["A5"].font = Font(bold=True, size=18, color="1F3864")
    ws["A6"] = f"Run ID: {run_id}"
    ws["A6"].font = Font(italic=True, color="666666")

    header_row = 8
    for col_idx, column_name in enumerate(invoice_df.columns, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=column_name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row in enumerate(invoice_df.itertuples(index=False), start=header_row + 1):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    for col_idx, column_name in enumerate(invoice_df.columns, start=1):
        max_len = max(
            [len(str(column_name))]
            + [len(str(v)) for v in invoice_df.iloc[:, col_idx - 1].tolist()]
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 45)

    amount_col_idx = _find_column_index(invoice_df, "Amount")
    if amount_col_idx:
        for row_idx in range(header_row + 1, header_row + 1 + len(invoice_df)):
            ws.cell(row=row_idx, column=amount_col_idx).number_format = "#,##0.00"

    ws.freeze_panes = f"A{header_row + 1}"

    summary_ws = wb.create_sheet(_SUMMARY_SHEET)
    _write_summary_sheet(summary_ws, summary_df)

    hidden_ws = wb.create_sheet(_RECONCILIATION_SHEET)
    for col_idx, column_name in enumerate(reconciliation_df.columns, start=1):
        hidden_ws.cell(row=1, column=col_idx, value=column_name)
    for row_idx, row in enumerate(reconciliation_df.itertuples(index=False), start=2):
        for col_idx, value in enumerate(row, start=1):
            hidden_ws.cell(row=row_idx, column=col_idx, value=value)
    hidden_ws.sheet_state = "hidden"

    wb.save(output_path)


def _write_summary_sheet(ws, summary_df: pd.DataFrame) -> None:
    ws["A1"] = "Property Summary"
    ws["A1"].font = Font(bold=True, size=16, color="1F3864")

    header_row = 3
    for col_idx, column_name in enumerate(summary_df.columns, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=column_name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row in enumerate(summary_df.itertuples(index=False), start=header_row + 1):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    for row_idx in range(header_row + 1, header_row + 1 + len(summary_df)):
        ws.cell(row=row_idx, column=2).number_format = "#,##0"
        ws.cell(row=row_idx, column=3).number_format = "#,##0.00"

    for col_idx, column_name in enumerate(summary_df.columns, start=1):
        max_len = max(
            [len(str(column_name))]
            + [len(str(v)) for v in summary_df.iloc[:, col_idx - 1].tolist()]
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 50)


def _write_invoice_pdf(output_path: Path, summary_df: pd.DataFrame, run_id: str) -> None:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("reportlab is required. Run: pip install reportlab") from exc

    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=letter,
        rightMargin=0.55 * inch,
        leftMargin=0.55 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
    )
    styles = getSampleStyleSheet()

    logo_cells = []
    for path, width in ((_REDPOINT_LOGO_PATH, 2.25 * inch), (_CREDIT_BOOST_LOGO_PATH, 1.8 * inch)):
        if path.exists():
            logo_cells.append(Image(str(path), width=width, height=0.75 * inch, kind="proportional"))
        else:
            logo_cells.append("")

    elements = [
        Table([logo_cells], colWidths=[2.75 * inch, 2.25 * inch], hAlign="LEFT"),
        Spacer(1, 0.2 * inch),
        Paragraph("Invoice", styles["Title"]),
        Paragraph(f"Run ID: {run_id}", styles["Normal"]),
        Spacer(1, 0.25 * inch),
    ]

    table_rows = [summary_df.columns.tolist()]
    for row in summary_df.itertuples(index=False):
        property_name, people_count, total_amount = row
        table_rows.append([property_name, people_count, f"${Decimal(str(total_amount)):.2f}"])

    table = Table(table_rows, colWidths=[3.8 * inch, 1.1 * inch, 1.4 * inch], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3864")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D9E2F3")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F8FB")]),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
    ]))
    elements.append(table)

    doc.build(elements)


def _build_property_summary(invoice_df: pd.DataFrame) -> pd.DataFrame:
    property_col = _find_source_column(invoice_df, {"propertyname", "property"})
    amount_col = _find_amount_column(invoice_df)

    rows = []
    for property_name, group in invoice_df.groupby(property_col, dropna=False):
        total_amount = sum(
            (parse_decimal(value, default=Decimal("0")) for value in group[amount_col]),
            Decimal("0"),
        )
        rows.append({
            "Property": str(property_name).strip(),
            "People Count": int(len(group)),
            "Total Amount": float(total_amount.quantize(Decimal("0.01"))),
        })

    rows.sort(key=lambda r: r["Property"].lower())
    total_people = sum(r["People Count"] for r in rows)
    total_amount = sum(Decimal(str(r["Total Amount"])) for r in rows)
    rows.append({
        "Property": "TOTAL",
        "People Count": total_people,
        "Total Amount": float(total_amount.quantize(Decimal("0.01"))),
    })
    return pd.DataFrame(rows, columns=["Property", "People Count", "Total Amount"])


def _find_source_column(df: pd.DataFrame, column_keys: set[str]) -> str:
    for col in df.columns:
        if _column_key(col) in column_keys:
            return col
    raise ValueError(f"Required column not found. Columns: {list(df.columns)}")


def _add_logo(ws, image_path: Path, anchor: str, *, width: int) -> None:
    if not image_path.exists():
        logger.warning("Logo not found for Redpoint invoice: %s", image_path)
        return

    image = OpenpyxlImage(str(image_path))
    if image.width:
        ratio = width / image.width
        image.width = width
        image.height = int(image.height * ratio)
    ws.add_image(image, anchor)


def _find_column_index(df: pd.DataFrame, column_name: str) -> Optional[int]:
    for idx, col in enumerate(df.columns, start=1):
        if str(col).strip().lower() == column_name.lower():
            return idx
    return None


def _resolve_base_price(base_price: Optional[Decimal]) -> Decimal:
    if base_price is not None:
        return Decimal(str(base_price)).quantize(Decimal("0.01"))

    try:
        from config import Config

        return Decimal(str(Config.CREDIT_BOOST_BASE_PRICE)).quantize(Decimal("0.01"))
    except Exception:
        return Decimal("6.50")


def _find_amount_column(df: pd.DataFrame) -> str:
    for col in df.columns:
        if _column_key(col) in _AMOUNT_COLUMN_CANDIDATES:
            return col
    raise ValueError(f"Boom statement is missing an Amount column. Columns: {list(df.columns)}")


def _column_key(column_name: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(column_name).strip().lower())
