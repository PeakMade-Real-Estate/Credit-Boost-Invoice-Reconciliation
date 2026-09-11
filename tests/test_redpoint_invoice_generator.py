from __future__ import annotations

import io
import zipfile
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
import pandas as pd

from app import create_app
from config import Config
from services.boom_invoice_parser import parse_boom_file
from services.redpoint_invoice_generator import generate_redpoint_invoice, generate_redpoint_invoice_package

FIXTURES = Path(__file__).parent / "fixtures"


def test_generate_redpoint_invoice_reprices_and_removes_boom_columns(tmp_path):
    source = tmp_path / "boom_statement.csv"
    pd.DataFrame(
        [
            {
                "ID": "6607745",
                "Item": "Boom fee",
                "Name": "Adam Eldridge",
                "Email": "adam@example.com",
                "Category": "partner_boom_report_ongoing_fee",
                "Description": "Boom fee for Adam Eldridge",
                "Template Name": "BoomReport",
                "Address": "123 Main St",
                "Property ID": "682348",
                "Property Name": "Beach Club",
                "Property Address": "123 Main St",
                "Property Group Name": "",
                "Unit ID": "",
                "Unit": "",
                "Amount": "-$0.91",
                "Transaction Type": "Invoice",
                "Created At": "2026-07-31T22:29:56.367+00:00",
                "Created At (Central)": "2026-07-31 5:29 PM",
                "Applicant Submitted Date": "",
                "Statement ID": "446699",
                "Bank Account Name": "",
                "Bank Account Last 4": "",
                "Subject Type": "Boom::ReportingAccount",
            }
        ]
    ).to_csv(source, index=False)

    output_path = generate_redpoint_invoice(
        str(source), str(tmp_path), "REC-TEST", base_price=Decimal("6.50")
    )

    assert output_path.endswith(".xlsx")

    wb = load_workbook(output_path)
    assert "Redpoint Invoice" in wb.sheetnames
    assert "Property Summary" in wb.sheetnames
    assert "Reconciliation Data" in wb.sheetnames
    assert wb["Reconciliation Data"].sheet_state == "hidden"
    assert len(wb["Redpoint Invoice"]._images) == 2

    df = pd.read_excel(output_path, sheet_name="Redpoint Invoice", header=7, dtype=str)
    assert df["Amount"].tolist() == ["6.50"]
    assert df.columns.tolist() == [
        "Name",
        "Email",
        "Address",
        "Property Name",
        "Property Address",
        "Property Group Name",
        "Unit ID",
        "Unit",
        "Amount",
        "Created At",
        "Created At (Central)",
        "Applicant Submitted Date",
    ]

    lines, summary = parse_boom_file(output_path, reporting_month="2026-07")
    assert summary.line_count == 1
    assert lines[0].transaction_id == "6607745"
    assert lines[0].boom_property_id == "Beach Club"
    assert lines[0].category == "partner_boom_report_ongoing_fee"
    assert lines[0].transaction_amount == Decimal("6.50")


def test_generate_redpoint_invoice_package_adds_summary_and_pdf(tmp_path):
    source = tmp_path / "boom_statement.csv"
    pd.DataFrame(
        [
            {"ID": "1", "Name": "A", "Property Name": "Beach Club", "Amount": "-$0.91"},
            {"ID": "2", "Name": "B", "Property Name": "Beach Club", "Amount": "-$0.91"},
            {"ID": "3", "Name": "C", "Property Name": "Campus Creek", "Amount": "-$0.91"},
        ]
    ).to_csv(source, index=False)

    xlsx_path, pdf_path = generate_redpoint_invoice_package(
        str(source), str(tmp_path), "REC-TEST", base_price=Decimal("6.50")
    )

    assert Path(pdf_path).exists()
    assert Path(pdf_path).stat().st_size > 0

    wb = load_workbook(xlsx_path)
    ws = wb["Property Summary"]
    assert ws["A4"].value == "Beach Club"
    assert ws["B4"].value == 2
    assert ws["C4"].value == 13.0
    assert ws["A6"].value == "TOTAL"
    assert ws["B6"].value == 3
    assert ws["C6"].value == 19.5


def test_redpoint_invoice_route_returns_workbook(tmp_path):
    class TestConfig(Config):
        TESTING = True
        UPLOAD_FOLDER = str(tmp_path / "uploads")
        OUTPUT_FOLDER = str(tmp_path / "outputs")
        DATABASE_PATH = str(tmp_path / "test.db")
        SECRET_KEY = "test-secret"

    source = tmp_path / "boom_statement.csv"
    pd.DataFrame(
        [
            {
                "ID": "6607745",
                "Item": "Boom fee",
                "Name": "Adam Eldridge",
                "Email": "adam@example.com",
                "Category": "partner_boom_report_ongoing_fee",
                "Description": "Boom fee for Adam Eldridge",
                "Template Name": "BoomReport",
                "Property ID": "682348",
                "Property Name": "Beach Club",
                "Amount": "-$0.91",
                "Transaction Type": "Invoice",
                "Statement ID": "446699",
                "Bank Account Name": "",
                "Bank Account Last 4": "",
                "Subject Type": "Boom::ReportingAccount",
            }
        ]
    ).to_csv(source, index=False)

    app = create_app(TestConfig)
    client = app.test_client()

    with open(source, "rb") as fh:
        response = client.post(
            "/redpoint-invoice",
            data={"statement_file": (fh, "boom_statement.csv")},
            content_type="multipart/form-data",
        )

    assert response.status_code == 200
    assert response.headers["Content-Type"] == "application/zip"
    assert "redpoint_invoice_package.zip" in response.headers["Content-Disposition"]

    with zipfile.ZipFile(io.BytesIO(response.data)) as zf:
        names = zf.namelist()
    assert any(name.endswith(".xlsx") for name in names)
    assert any(name.endswith(".pdf") for name in names)
